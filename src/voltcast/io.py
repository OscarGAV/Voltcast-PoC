"""Lectura del Excel de origen y lectura/escritura de las capas Medallion (Parquet)."""
from __future__ import annotations

import hashlib
import json
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import pandas as pd

from . import config

DATE_COL = "Date"
VOLTS_COL = "[Volts]"
KA_COL = "[kA]"
LINEAGE_COLS = ["electrolizador", "_source_file", "_ingested_at"]


class SchemaError(ValueError):
    """El archivo no tiene la estructura esperada (US001-E03, US002-E02, US003-E02)."""


# --------------------------------------------------------------------------- #
# Excel de origen
# --------------------------------------------------------------------------- #
def _header_to_text(value) -> str:
    """Nombre de columna a texto. Los números enteros (1.0) quedan como '1'."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_excel_sheet(path: Path, sheet: str) -> pd.DataFrame:
    """Lee una hoja con openpyxl en modo read_only, sin transformar valores.

    Solo se convierten los encabezados a texto y se descartan las filas completamente vacías.
    """
    path = Path(path)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise SchemaError(f"Formato no soportado: '{path.suffix}'. Se espera un archivo Excel (.xlsx).")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise SchemaError(f"La hoja '{sheet}' no existe en {path.name}. Hojas disponibles: {wb.sheetnames}")
        rows = wb[sheet].iter_rows(values_only=True)
        header = [_header_to_text(h) for h in next(rows)]
        data = [r for r in rows if any(v is not None for v in r)]
    finally:
        wb.close()
    df = pd.DataFrame(data, columns=header)
    # Columnas sin encabezado y completamente vacías (bordes de la hoja)
    empty_unnamed = [c for c in df.columns if c == "" and df[c].isna().all()]
    return df.drop(columns=empty_unnamed)


@dataclass
class SheetSchema:
    date_col: str
    volts_col: str
    ka_col: str
    cell_cols: list[str]
    other_cols: list[str] = field(default_factory=list)

    @property
    def n_cells(self) -> int:
        return len(self.cell_cols)


def detect_schema(df: pd.DataFrame, sheet: str = "") -> SheetSchema:
    """Identifica la fecha, [Volts], [kA] y las columnas de celdas (US002/US003).

    Las celdas se detectan por encabezado numérico entero; se cuentan, no se asume 181.
    Lanza SchemaError con un mensaje descriptivo si falta algo.
    """
    where = f" en la hoja '{sheet}'" if sheet else ""
    cols = list(df.columns)
    if DATE_COL not in cols:
        raise SchemaError(f"No se encontró la columna de fecha '{DATE_COL}'{where}. Verifique la estructura del archivo.")
    missing = [c for c in (VOLTS_COL, KA_COL) if c not in cols]
    if missing:
        raise SchemaError(f"No se encontraron las columnas {missing}{where} (voltaje total / corriente).")
    cell_cols = [c for c in cols if c.isdigit()]
    if not cell_cols:
        raise SchemaError(f"No se encontraron columnas de celdas eléctricas (encabezados numéricos){where}.")
    non_numeric = [c for c in cell_cols + [VOLTS_COL, KA_COL] if not pd.api.types.is_numeric_dtype(df[c])]
    if non_numeric:
        raise SchemaError(f"Columnas de voltaje/corriente con valores no numéricos{where}: {non_numeric[:10]}")
    if not pd.api.types.is_datetime64_any_dtype(df[DATE_COL]):
        raise SchemaError(f"La columna '{DATE_COL}'{where} no contiene fechas.")
    other = [c for c in cols if c not in {DATE_COL, VOLTS_COL, KA_COL, *cell_cols}]
    return SheetSchema(DATE_COL, VOLTS_COL, KA_COL, cell_cols, other)


def file_sha256(path: Path, chunk: int = 1 << 20) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        while block := f.read(chunk):
            h.update(block)
    return h.hexdigest()


# --------------------------------------------------------------------------- #
# Capas Medallion
# --------------------------------------------------------------------------- #
LAYER_DIRS = {"bronze": config.BRONZE_DIR, "silver": config.SILVER_DIR, "gold": config.GOLD_DIR}


def layer_path(layer: str, name: str) -> Path:
    if layer not in LAYER_DIRS:
        raise ValueError(f"Capa desconocida '{layer}'. Opciones: {list(LAYER_DIRS)}")
    return LAYER_DIRS[layer] / name


def write_layer(df: pd.DataFrame, layer: str, name: str, source: Path | str, source_hash: str | None = None,
                extra: dict | None = None) -> Path:
    """Escribe ``df`` como Parquet en la capa y registra su entrada en ``_metadata.json``."""
    path = layer_path(layer, name)
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_parquet(path, engine="pyarrow", index=False)

    source = Path(source)
    if source_hash is None and source.is_file():
        source_hash = file_sha256(source)
    entry = {
        "executed_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "rows": int(len(df)),
        "columns": int(df.shape[1]),
        "source_file": source.name,
        "source_sha256": source_hash,
        **(extra or {}),
    }
    meta_path = path.parent / "_metadata.json"
    meta = json.loads(meta_path.read_text(encoding="utf-8")) if meta_path.exists() else {}
    meta[name] = entry
    meta_path.write_text(json.dumps(meta, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    return path


def read_layer(layer: str, name: str) -> pd.DataFrame:
    path = layer_path(layer, name)
    if not path.exists():
        raise FileNotFoundError(f"No existe {path}. Ejecute primero el notebook que genera la capa '{layer}'.")
    return pd.read_parquet(path, engine="pyarrow")


def read_metadata(layer: str) -> dict:
    meta_path = layer_path(layer, "_metadata.json")
    return json.loads(meta_path.read_text(encoding="utf-8")) if meta_path.exists() else {}


# --------------------------------------------------------------------------- #
# Bronze
# --------------------------------------------------------------------------- #
def bronze_name(electrolizador: str) -> str:
    return f"voltajes_EL_{electrolizador}.parquet"


def read_bronze(electrolizador: str) -> pd.DataFrame:
    return read_layer("bronze", bronze_name(electrolizador))


def schema_to_dict(schema: SheetSchema) -> dict:
    d = asdict(schema)
    d["n_cells"] = schema.n_cells
    return d
